#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
주식 테마 분류 + 테마 뉴스속보 수집기
====================================

사용법
------
  python3 collector.py --source sample            # 샘플 데이터로 실행 (기본)
  python3 collector.py --source rss               # 실제 RSS 수집
  python3 collector.py --source rss --hours 12    # 최근 12시간만
  python3 collector.py --source sample --xlsx     # 엑셀 리포트까지 생성

출력
----
  data/news_data.json   : 대시보드(dashboard.html)가 읽는 데이터
  data/theme_report.xlsx: 엑셀 리포트 (--xlsx 옵션)

의존성: 표준 라이브러리만으로 동작. --xlsx 사용 시 openpyxl 필요.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import random
import re
import sys
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
KST = timezone(timedelta(hours=9))

# ---------------------------------------------------------------------------
# 1. 뉴스 소스 설정
# ---------------------------------------------------------------------------
# Google News RSS는 검색어 기반으로 국내/해외 매체를 한 번에 긁을 수 있어 기본값으로 둔다.
# 사내 정책상 특정 매체 RSS만 써야 한다면 아래 리스트만 갈아끼우면 된다.
GOOGLE_NEWS_RSS = (
    "https://news.google.com/rss/search?q={q}&hl={hl}&gl={gl}&ceid={ceid}"
)

RSS_SOURCES = [
    # (라벨, URL, 시장)
    ("연합인포맥스-증시", "https://news.einfomax.co.kr/rss/S1N2.xml", "KR"),
    ("한국경제-증권", "https://www.hankyung.com/feed/finance", "KR"),
    ("매일경제-증권", "https://www.mk.co.kr/rss/50200011/", "KR"),
]

# 테마별 검색어 기반 수집(선택). --source rss 실행 시 테마 키워드 상위 3개로 질의한다.
QUERY_PER_THEME = 3


def _fetch(url: str, timeout: int = 10) -> bytes:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (compatible; ThemeNewsBot/1.0)",
            "Accept": "application/rss+xml, application/xml, text/xml, */*",
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def _clean(text: str) -> str:
    text = re.sub(r"<[^>]+>", " ", text or "")
    text = html.unescape(text)
    return re.sub(r"\s+", " ", text).strip()


def _parse_rss(raw: bytes, source_label: str, market: str) -> list[dict]:
    """표준 RSS 2.0 / Atom 모두 대응하는 최소 파서."""
    items: list[dict] = []
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return items

    nodes = root.findall(".//item")
    atom_ns = "{http://www.w3.org/2005/Atom}"
    if not nodes:
        nodes = root.findall(f".//{atom_ns}entry")

    for node in nodes:
        def pick(*tags):
            for t in tags:
                el = node.find(t)
                if el is not None:
                    if el.text:
                        return el.text
                    if el.get("href"):
                        return el.get("href")
            return ""

        title = _clean(pick("title", f"{atom_ns}title"))
        if not title:
            continue
        link = _clean(pick("link", f"{atom_ns}link"))
        desc = _clean(pick("description", "summary", f"{atom_ns}summary"))[:400]
        pub_raw = pick("pubDate", "published", f"{atom_ns}published", "updated")

        published = None
        if pub_raw:
            try:
                published = parsedate_to_datetime(pub_raw)
            except (TypeError, ValueError):
                try:
                    published = datetime.fromisoformat(pub_raw.replace("Z", "+00:00"))
                except ValueError:
                    published = None
        if published is None:
            published = datetime.now(timezone.utc)
        if published.tzinfo is None:
            published = published.replace(tzinfo=timezone.utc)

        items.append(
            {
                "title": title,
                "summary": desc,
                "url": link,
                "source": source_label,
                "market": market,
                "published": published.astimezone(KST).isoformat(),
            }
        )
    return items


def collect_rss(themes: list[dict], hours: int, search_limit: int = 24) -> list[dict]:
    """설정된 RSS + 테마 키워드 검색으로 뉴스를 모은다.

    테마가 늘어나면 검색 요청도 같이 늘어난다(테마 42개 × 3 = 126회). 그대로 두면
    화면에 아무 표시 없이 수십 분이 걸리고, 구글이 중간에 막으면 조용히 실패한다.
    그래서 (a) 진행 상황을 한 줄씩 찍고 (b) 검색어 중복을 없애고 (c) search_limit 로
    상한을 둔다. 검색을 아예 끄려면 search_limit=0.
    """
    collected: list[dict] = []
    ok_cnt = fail_cnt = 0

    print(f"  · 매체 RSS {len(RSS_SOURCES)}곳")
    for label, url, market in RSS_SOURCES:
        try:
            got = _parse_rss(_fetch(url), label, market)
            collected += got
            ok_cnt += 1
            print(f"    [OK] {label} {len(got)}건")
        except Exception as exc:  # 개별 소스 실패가 전체를 막지 않도록
            fail_cnt += 1
            print(f"    [실패] {label} — {type(exc).__name__}: {str(exc)[:60]}")

    # 테마별 상위 키워드를 모아 중복을 제거한다. 'AI' 처럼 여러 테마가 공유하는
    # 키워드로 같은 검색을 몇 번씩 날릴 이유가 없다.
    queries: list[tuple[str, str]] = []
    seen_kw: set[str] = set()
    for theme in themes:
        for k in theme["keywords"][:QUERY_PER_THEME]:
            key = k["k"].lower()
            if key not in seen_kw:
                seen_kw.add(key)
                queries.append((k["k"], theme["market"]))
    if search_limit > 0:
        queries = queries[:search_limit]

        print(f"  · 구글뉴스 검색 {len(queries)}개 (--search-limit 로 조절)")
        for i, (kw, market) in enumerate(queries, start=1):
            hl, gl, ceid = ("ko", "KR", "KR:ko") if market == "KR" else ("en-US", "US", "US:en")
            q = urllib.parse.quote(f'"{kw}" 주가' if market == "KR" else f'"{kw}" stock')
            url = GOOGLE_NEWS_RSS.format(q=q, hl=hl, gl=gl, ceid=ceid)
            try:
                got = _parse_rss(_fetch(url), f"GoogleNews:{kw}", market)
                collected += got
                ok_cnt += 1
                print(f"    [{i}/{len(queries)}] {kw} {len(got)}건")
            except Exception as exc:
                fail_cnt += 1
                print(f"    [{i}/{len(queries)}] {kw} 실패 — {type(exc).__name__}: {str(exc)[:50]}")

    print(f"  · 요청 성공 {ok_cnt} / 실패 {fail_cnt} · 원본 {len(collected)}건")
    if not collected:
        print("  [!] 한 건도 못 받았습니다. 인터넷 연결이나 회사 방화벽을 확인해 주세요.")

    cutoff = datetime.now(KST) - timedelta(hours=hours)
    fresh = [n for n in collected if datetime.fromisoformat(n["published"]) >= cutoff]
    if collected and not fresh:
        print(f"  [!] 받은 {len(collected)}건이 전부 {hours}시간보다 오래된 기사입니다."
              f" --hours 를 늘려보세요.")
    return dedupe(fresh)


def dedupe(news: list[dict]) -> list[dict]:
    """제목 정규화 기준 중복 제거 (같은 기사 재송고/전재 대응)."""
    seen: dict[str, dict] = {}
    for n in news:
        key = re.sub(r"[^0-9a-zA-Z가-힣]", "", n["title"])[:60]
        if key not in seen:
            seen[key] = n
    return list(seen.values())


# ---------------------------------------------------------------------------
# 2. 테마 분류 엔진
# ---------------------------------------------------------------------------
URGENCY_PATTERNS = [
    (3, r"속보|긴급|장중|사상 최대|급등|급락|상한가|하한가|breaking|surge|plunge|halted"),
    (2, r"단독|잠정|첫|최초|돌파|신고가|계약 체결|수주|승인|허가|beats|soars|record"),
    (1, r"전망|기대|검토|추진|계획|outlook|expects|plans"),
]

POSITIVE = r"급등|상승|호재|수주|계약|수출|흑자|증가|확대|돌파|신고가|승인|허가|최대|성장|호실적|목표가 상향|beats|surge|soar|record|approval|upgrade|wins"
NEGATIVE = r"급락|하락|악재|취소|철회|적자|감소|축소|리콜|제재|규제|소송|우려|부진|목표가 하향|miss|plunge|halt|recall|probe|downgrade|lawsuit"


def _name_hit(name: str, text: str) -> bool:
    """종목명 매칭에는 단어 경계를 요구한다.

    한글에는 \\b 가 통하지 않아 '인텔'이 '인텔리안테크'에, '우진'이 '우진산전'에
    걸리는 오탐이 난다. 앞뒤에 한글/영문이 붙어 있으면 다른 단어로 보고 버린다.
    """
    pat = rf"(?<![가-힣A-Za-z]){re.escape(name)}(?![가-힣A-Za-z])"
    return re.search(pat, text, re.IGNORECASE) is not None


def score_theme(text: str, theme: dict) -> tuple[float, list[str], bool]:
    """키워드 가중합으로 테마 적합도 계산. (점수, 매칭키워드, 종목명직접언급여부)"""
    score = 0.0
    hits: list[str] = []
    named = False
    low = text.lower()
    for kw in theme["keywords"]:
        k = kw["k"]
        if k.lower() in low:
            score += kw["w"]
            hits.append(k)
    # 종목명 직접 언급은 강한 신호. alias 는 국내 매체가 쓰는 한글 표기·약칭.
    for st in theme["stocks"]:
        names = [st["name"]] + st.get("alias", [])
        bonus = {"leader": 3.5, "core": 3.0}.get(st["role"], 3.0)
        matched_name = next((n for n in names if n != "-" and _name_hit(n, text)), None)
        if matched_name:
            score += bonus
            hits.append(matched_name)
            named = True
        elif st["code"] not in ("-", "") and re.search(rf"\b{re.escape(st['code'])}\b", text):
            score += 1.5
            hits.append(st["code"])
            named = True
    return score, hits, named


def classify(news: dict, themes: list[dict], min_score: float = 3.0) -> dict:
    text = f"{news['title']} {news.get('summary', '')}"
    results = []
    for theme in themes:
        s, hits, named = score_theme(text, theme)
        # 시장이 다르면 감점(국내 기사가 미국 테마에 끌려가는 것 방지).
        # 단 (a) 종목명이 직접 언급됐거나 (b) 시장 무관 테마(ETF·제도·리포트 등)면 감점하지 않는다.
        # 종목명은 시장 판정보다 확실한 신호다 — '루멘텀 실적' 기사에 美 표식이 없다고
        # 미국 테마에서 밀어내면 오히려 손해다.
        mismatch = theme["market"] != news.get("market", theme["market"])
        if mismatch and not named and not theme.get("cross_market"):
            s *= 0.6
        if s >= min_score:
            results.append({"theme_id": theme["id"], "theme": theme["name"],
                            "market": theme["market"], "score": round(s, 1),
                            "matched": sorted(set(hits))[:6]})
    results.sort(key=lambda r: -r["score"])

    urgency = 0
    for level, pat in URGENCY_PATTERNS:
        if re.search(pat, text, re.IGNORECASE):
            urgency = max(urgency, level)

    pos = len(re.findall(POSITIVE, text, re.IGNORECASE))
    neg = len(re.findall(NEGATIVE, text, re.IGNORECASE))
    sentiment = "positive" if pos > neg else "negative" if neg > pos else "neutral"

    top = results[0] if results else None
    news = dict(news)
    news.update(
        {
            "themes": results[:3],
            "primary_theme": top["theme_id"] if top else None,
            "primary_theme_name": top["theme"] if top else "미분류",
            "confidence": min(100, round((top["score"] / 12) * 100)) if top else 0,
            "urgency": urgency,
            "sentiment": sentiment,
            "related_stocks": next(
                (t["stocks"][:4] for t in themes if top and t["id"] == top["theme_id"]), []
            ),
        }
    )
    return news


def build_theme_stats(classified: list[dict], themes: list[dict]) -> list[dict]:
    """테마별 뉴스량·속보량·심리 집계 = 대시보드의 '테마 열기' 지표."""
    stats = []
    now = datetime.now(KST)
    for theme in themes:
        rel = [n for n in classified if any(t["theme_id"] == theme["id"] for t in n["themes"])]
        if not rel:
            stats.append({**{k: theme[k] for k in ("id", "name", "market", "sector", "summary")},
                          "stocks": theme["stocks"], "news_count": 0, "breaking_count": 0,
                          "heat": 0, "sentiment_score": 0, "last_news": None})
            continue
        breaking = sum(1 for n in rel if n["urgency"] >= 3)
        pos = sum(1 for n in rel if n["sentiment"] == "positive")
        neg = sum(1 for n in rel if n["sentiment"] == "negative")
        # 열기 = 기사 수 + 속보 가중 + 최신성 가중
        recency = sum(
            max(0.0, 1 - (now - datetime.fromisoformat(n["published"])).total_seconds() / 86400)
            for n in rel
        )
        heat = round(len(rel) * 1.0 + breaking * 2.0 + recency * 1.5, 1)
        stats.append(
            {
                **{k: theme[k] for k in ("id", "name", "market", "sector", "summary")},
                "stocks": theme["stocks"],
                "news_count": len(rel),
                "breaking_count": breaking,
                "heat": heat,
                "sentiment_score": round((pos - neg) / len(rel) * 100),
                "last_news": max(n["published"] for n in rel),
            }
        )
    stats.sort(key=lambda s: -s["heat"])
    return stats


# ---------------------------------------------------------------------------
# 3. 샘플 데이터 생성기 (구조 검증 / 오프라인 데모용)
# ---------------------------------------------------------------------------
SAMPLE_TEMPLATES = {
    "kr-hbm": [
        ("[속보] SK하이닉스, HBM4 12단 엔비디아 퀄테스트 통과…내년 물량 선점", "SK하이닉스가 차세대 고대역폭 메모리 HBM4 12단 제품의 고객사 품질 검증을 마쳤다고 밝혔다. 업계는 내년 공급 물량의 절반 이상을 확보한 것으로 본다."),
        ("삼성전자 D램 가격 3개월 연속 상승…메모리 업사이클 진입", "DDR5 고정거래가격이 전월 대비 8% 오르며 상승세를 이어갔다."),
        ("한미반도체, TC본더 수주 잔고 사상 최대…하이브리드 본딩 장비도 공급", "HBM 후공정 장비 수요가 이어지며 수주 잔고가 분기 기준 최대치를 기록했다."),
    ],
    "kr-battery": [
        ("LG에너지솔루션, 북미 ESS 배터리 5조원 규모 장기 공급 계약", "LFP 기반 ESS 셀을 2031년까지 공급하는 계약을 체결했다고 공시했다."),
        ("전고체 배터리 파일럿 라인 가동 임박…소재주 동반 급등", "삼성SDI가 전고체 시험생산 라인 가동 시점을 앞당기면서 관련 소재 기업 주가가 강세를 보였다."),
        ("양극재 판가 하락에 에코프로비엠 목표가 하향", "리튬 가격 약세가 판가에 반영되며 증권사들이 목표주가를 낮췄다."),
    ],
    "kr-nuclear": [
        ("[속보] 두산에너빌리티, 체코 원전 주기기 공급 계약 최종 서명", "체코 두코바니 신규 원전 2기에 대한 주기기 공급 계약이 최종 체결됐다."),
        ("정부, SMR 실증사업 예산 확대…한전기술 설계 참여", "소형모듈원자로 표준설계 인가 절차가 내년 상반기 마무리될 전망이다."),
    ],
    "kr-defense": [
        ("한화에어로스페이스, 중동 K9 자주포 2조원 수주", "사우디아라비아와 K9 자주포 및 탄약 패키지 공급 계약을 체결했다."),
        ("LIG넥스원 천궁-II 추가 수출 협상 진행…방산 수출 사상 최대 전망", "올해 방산 수출액이 사상 처음 250억달러를 넘어설 것이라는 전망이 나왔다."),
    ],
    "kr-power": [
        ("HD현대일렉트릭, 미국 변압기 공장 증설…전력망 교체 수요 대응", "북미 초고압 변압기 수주 잔고가 3년치를 넘어섰다."),
        ("데이터센터 전력 수요 급증에 전력기기 3사 목표가 일제히 상향", "증권가는 송배전 설비 슈퍼사이클이 2030년까지 이어질 것으로 봤다."),
    ],
    "kr-bio": [
        ("알테오젠, 글로벌 빅파마와 피하주사 기술수출 계약…마일스톤 1조원", "총 계약 규모 1조2000억원의 라이선스아웃 계약을 체결했다고 밝혔다."),
        ("삼성바이오로직스 5공장 가동…CDMO 수주 잔고 확대", "생산능력 기준 세계 1위 지위를 굳혔다는 평가가 나온다."),
        ("셀트리온 바이오시밀러 미국 FDA 품목허가 획득", "자가면역질환 치료제의 미국 판매 허가를 받았다."),
    ],
    "kr-robot": [
        ("[속보] 레인보우로보틱스, 휴머노이드 양산 라인 착공", "삼성전자와 공동 개발한 휴머노이드 로봇의 양산 준비에 들어갔다."),
        ("감속기 국산화 성공…에스비비테크 공급 계약 체결", "하모닉 감속기 국산화로 로봇 부품 밸류체인이 확대되고 있다."),
    ],
    "kr-crypto": [
        ("원화 스테이블코인 법제화 논의 재개…관련주 급등", "국회 정무위가 디지털자산 2단계 법안 심사에 착수했다."),
        ("비트코인 신고가 경신에 가상자산 관련주 동반 강세", "거래대금 증가로 거래소 수수료 수익이 늘어날 것이라는 기대가 반영됐다."),
    ],
    "kr-ent": [
        ("하이브 소속 아티스트 월드투어 발표…엔터주 강세", "북미·유럽 40개 도시 투어가 확정되며 공연 매출 기대가 커졌다."),
        ("K팝 앨범 수출 역대 최대…엔터 3사 실적 눈높이 상향", "상반기 음반 수출액이 전년 대비 30% 늘었다."),
    ],
    "kr-beauty": [
        ("실리콘투, 미국 아마존 뷰티 매출 급증…분기 사상 최대 실적", "K뷰티 인디 브랜드 유통 물량이 크게 늘었다."),
        ("화장품 수출 3개월 연속 증가…ODM 가동률 90% 육박", "미국·일본향 수출이 성장을 견인했다."),
    ],
    "kr-ship": [
        ("한화오션, LNG운반선 6척 수주…선가 지수 재차 상승", "클락슨 신조선가지수가 연중 최고치를 기록했다."),
        ("미국 해군 MRO 사업 확대 논의…국내 조선사 수혜 기대", "한미 조선 협력 논의가 구체화되고 있다."),
    ],
    "kr-value": [
        ("금융지주 자사주 소각 확대 발표…밸류업 기대 재점화", "주요 금융지주가 연간 총주주환원율 50% 목표를 제시했다."),
        ("상법 개정안 국회 통과…지배구조 개선 기대", "이사 충실의무 확대 조항이 포함됐다."),
    ],
    "kr-ai-sw": [
        ("정부 소버린 AI 프로젝트 2차 사업자 선정…네이버 컨소시엄 참여", "국산 LLM 개발에 대규모 예산이 투입된다."),
        ("국내 데이터센터 투자 확대…IDC 관련주 주목", "AI 학습 수요로 국내 코로케이션 수요가 늘고 있다."),
    ],
    "kr-auto": [
        ("미국 자동차 관세 인하 협상 진전…완성차 관련주 반등", "관세율 조정 가능성이 제기되며 투자심리가 개선됐다."),
        ("현대차 로보택시 실증 확대…자율주행 부품주 강세", "미국 내 자율주행 실증 지역을 확대한다고 밝혔다."),
    ],
    "kr-game": [
        ("크래프톤 신작 글로벌 출시…스팀 동시접속 신기록", "출시 첫 주 판매량이 시장 기대치를 웃돌았다."),
        ("중국 판호 발급 재개…게임주 일제히 상승", "국내 게임 3종이 외자판호를 받았다."),
    ],
    "kr-hydrogen": [
        ("청정수소 발전 입찰 낙찰 결과 발표…두산퓨얼셀 물량 확보", "연료전지 발전 물량이 전년 대비 두 배로 늘었다."),
        ("해상풍력 고정가격 입찰 확대…씨에스윈드 수주 기대", "정부가 해상풍력 보급 목표를 상향했다."),
    ],
    "kr-foundry": [
        ("삼성전자 2나노 파운드리 대형 고객 확보…소부장 수혜 기대", "테일러 공장 가동률이 빠르게 올라올 전망이다."),
        ("EUV 소재 국산화 진전…동진쎄미켐 공급 확대", "포토레지스트 국산화 비중이 높아지고 있다."),
    ],
    "kr-space": [
        ("누리호 5차 발사 성공…우주항공 관련주 강세", "실용위성을 목표 궤도에 안착시켰다."),
        ("저궤도 위성통신 사업 예타 통과", "2030년까지 위성 다수를 발사하는 계획이 확정됐다."),
    ],
    "us-ai-infra": [
        ("Nvidia beats estimates as Blackwell shipments surge; datacenter revenue hits record", "The company guided above consensus, citing sustained hyperscaler capex."),
        ("Broadcom wins new custom AI chip customer, backlog expands", "Management said AI revenue would roughly double next year."),
        ("Hyperscaler capex guidance raised again, lifting AI supply chain", "Combined 2027 capex plans now exceed prior estimates by a wide margin."),
    ],
    "us-power": [
        ("GE Vernova sold out on gas turbines through 2029 as datacenter power demand climbs", "Order backlog reached a record on AI-driven electricity demand."),
        ("Utility signs long-term PPA with hyperscaler for nuclear power", "The deal covers baseload supply to a new datacenter campus."),
    ],
    "us-nuclear": [
        ("NRC approves SMR design certification; Oklo shares surge", "The approval clears a key regulatory hurdle for first deployment."),
        ("Uranium prices hit multi-year high on supply disruption", "Cameco flagged production risk at a key mine."),
    ],
    "us-obesity": [
        ("Oral GLP-1 phase 3 data beats expectations on weight loss", "Eli Lilly reported strong efficacy with a manageable safety profile."),
        ("Novo Nordisk cuts guidance as US pricing pressure builds", "Shares plunged on the revised outlook."),
    ],
    "us-quantum": [
        ("IonQ reports error correction milestone, shares soar", "The company said logical qubit fidelity improved substantially."),
        ("Government awards quantum computing contracts to multiple vendors", "Funding targets national security applications."),
    ],
    "us-crypto": [
        ("Bitcoin ETF flows hit record as stablecoin rules take effect", "Coinbase and Circle rallied on regulatory clarity."),
        ("Stablecoin issuer reserves under SEC review", "Shares fell on the probe headline."),
    ],
    "us-space": [
        ("Rocket Lab wins Pentagon hypersonic test contract", "The award expands its defense revenue base."),
        ("Palantir lands expanded defense software deal", "The contract covers multi-year AI deployment."),
    ],
    "us-robotics": [
        ("Tesla expands robotaxi service area; Optimus production timeline updated", "Management reiterated humanoid volume targets."),
        ("Warehouse automation orders accelerate, Symbotic backlog grows", "Retail customers expanded deployment plans."),
    ],
    "us-cyber": [
        ("Major ransomware breach hits enterprise software vendor; cybersecurity names rally", "CrowdStrike and Palo Alto rose on expected budget shifts."),
        ("Zero trust adoption drives platform consolidation", "Analysts upgraded the sector on durable spending."),
    ],
    "us-semicap": [
        ("ASML bookings beat on EUV demand from foundry capex", "Management raised full-year guidance."),
        ("New export control rules tighten equipment shipments to China", "Semicap shares fell on the revised restrictions."),
    ],
    "us-cloud": [
        ("Azure growth reaccelerates on AI workloads; Copilot seats expand", "Microsoft cloud revenue topped estimates."),
        ("AWS backlog grows as enterprises commit to multi-year AI contracts", "Amazon flagged capacity constraints easing."),
    ],
    "us-macro": [
        ("Fed signals rate cut at next FOMC as CPI cools", "Treasury yields fell sharply following the release."),
        ("New tariff announcement rattles markets; risk assets plunge", "Equity futures dropped on the trade headline."),
    ],
}

SAMPLE_SOURCES_KR = ["연합인포맥스", "한국경제", "매일경제", "머니투데이", "이데일리", "서울경제"]
SAMPLE_SOURCES_US = ["Reuters", "Bloomberg", "CNBC", "WSJ", "Barron's"]


def generate_sample(themes: list[dict], hours: int, seed: int = 20260813) -> list[dict]:
    rng = random.Random(seed)
    now = datetime.now(KST)
    out: list[dict] = []
    by_id = {t["id"]: t for t in themes}
    for tid, rows in SAMPLE_TEMPLATES.items():
        theme = by_id.get(tid)
        if not theme:
            continue
        for title, summary in rows:
            mins = rng.randint(3, max(10, hours * 60))
            out.append(
                {
                    "title": title,
                    "summary": summary,
                    "url": f"https://example.com/news/{tid}-{abs(hash(title)) % 100000}",
                    "source": rng.choice(SAMPLE_SOURCES_KR if theme["market"] == "KR" else SAMPLE_SOURCES_US),
                    "market": theme["market"],
                    "published": (now - timedelta(minutes=mins)).isoformat(),
                }
            )
    rng.shuffle(out)
    return out


def load_file(path: str, default_market: str = "KR") -> list[dict]:
    """텍스트 덤프에서 뉴스를 읽는다.

    한 줄에 하나씩, 아래 두 형식을 지원한다.
        제목 ||| 출처 ||| 발행일시
        제목 ||| 발행일시
    '#'으로 시작하는 줄과 빈 줄은 무시. RSS를 못 쓰는 망 환경에서
    사내 뉴스 덤프·크롤링 결과를 그대로 밀어 넣을 때 쓴다.
    """
    out: list[dict] = []
    now = datetime.now(KST)
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = [p.strip() for p in line.split("|||")]
            title = parts[0]
            source = parts[1] if len(parts) >= 3 else "file"
            date_raw = parts[-1] if len(parts) >= 2 else ""
            published = now
            if date_raw:
                try:
                    published = parsedate_to_datetime(date_raw)
                except (TypeError, ValueError):
                    try:
                        published = datetime.fromisoformat(date_raw)
                    except ValueError:
                        published = now
            if published.tzinfo is None:
                published = published.replace(tzinfo=KST)
            # 제목에 美/뉴욕/나스닥 표식이 있으면 미국물로 본다
            market = "US" if re.search(r"美|뉴욕|나스닥|월가|NYSE|Nasdaq", title) else default_market
            out.append({"title": title, "summary": "", "url": "",
                        "source": source, "market": market,
                        "published": published.astimezone(KST).isoformat()})
    return dedupe(out)


# ---------------------------------------------------------------------------
# 4. 출력
# ---------------------------------------------------------------------------
def write_json(payload: dict) -> str:
    os.makedirs(DATA_DIR, exist_ok=True)
    path = os.path.join(DATA_DIR, "news_data.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return path


def write_xlsx(payload: dict) -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  [warn] openpyxl 미설치 - 엑셀 생략 (pip install openpyxl)", file=sys.stderr)
        return None

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(color="FFFFFF", bold=True)

    def style_header(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "테마요약"
    ws.append(["순위", "시장", "테마", "섹터", "열기", "뉴스수", "속보수", "심리(-100~100)", "대장주", "최근뉴스"])
    for i, t in enumerate(payload["themes"], start=1):
        leaders = ", ".join(s["name"] for s in t["stocks"] if s["role"] == "leader")
        ws.append([i, t["market"], t["name"], t["sector"], t["heat"], t["news_count"],
                   t["breaking_count"], t["sentiment_score"], leaders,
                   (t["last_news"] or "")[:16].replace("T", " ")])
    style_header(ws, [6, 8, 34, 14, 8, 8, 8, 16, 34, 18])

    ws2 = wb.create_sheet("뉴스")
    ws2.append(["시각", "시장", "긴급도", "심리", "대표테마", "제목", "매칭키워드", "출처", "링크"])
    for n in payload["news"]:
        matched = ", ".join(n["themes"][0]["matched"]) if n["themes"] else ""
        ws2.append([n["published"][:16].replace("T", " "), n["market"], n["urgency"],
                    n["sentiment"], n["primary_theme_name"], n["title"], matched,
                    n["source"], n["url"]])
    style_header(ws2, [18, 8, 9, 10, 28, 70, 34, 16, 40])

    ws3 = wb.create_sheet("테마종목")
    ws3.append(["시장", "테마", "종목명", "코드", "구분"])
    for t in payload["themes"]:
        for s in t["stocks"]:
            ws3.append([t["market"], t["name"], s["name"], s["code"], s["role"]])
    style_header(ws3, [8, 34, 26, 12, 10])

    os.makedirs(DATA_DIR, exist_ok=True)
    path = os.path.join(DATA_DIR, "theme_report.xlsx")
    wb.save(path)
    return path


def inject_into_dashboard(payload: dict) -> str | None:
    """dashboard.html 안의 데이터 블록을 최신 데이터로 교체."""
    path = os.path.join(BASE_DIR, "dashboard.html")
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        src = f.read()
    start = src.find("/*DATA_START*/")
    end = src.find("/*DATA_END*/")
    if start == -1 or end == -1:
        return None
    new = (src[: start + len("/*DATA_START*/")]
           + "\nconst EMBEDDED_DATA = "
           + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
           + ";\n"
           + src[end:])
    with open(path, "w", encoding="utf-8") as f:
        f.write(new)
    return path


class _Tee:
    """화면에 찍는 내용을 로그 파일에도 같이 남긴다.

    윈도우 배치 창은 끝나면 닫혀서, 문제가 생겨도 메시지를 다시 볼 수 없다.
    파이썬 쪽에서 남겨두면 창이 닫혀도 원인을 찾을 수 있다.
    """

    def __init__(self, stream, path: str):
        self.stream = stream
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            self.file = open(path, "w", encoding="utf-8-sig")   # 메모장 호환
        except Exception:
            self.file = None

    def write(self, s):
        self.stream.write(s)
        if self.file:
            try:
                self.file.write(s)
                self.file.flush()      # 중간에 죽어도 거기까지는 남게
            except Exception:
                pass
        return len(s)

    def flush(self):
        self.stream.flush()


def _safe_console() -> None:
    """출력하다 죽지 않게 한다.

    콘솔이 표현 못 하는 글자가 하나라도 있으면 print 가 UnicodeEncodeError 를 내며
    프로그램이 통째로 죽는다 — 30초마다 도는 갱신이 조용히 멈추는 원인이 된다.
    인코딩은 콘솔 것을 그대로 두고(바꾸면 한글이 깨진다) 실패만 '?' 로 넘긴다.
    """
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(errors="replace")     # py3.7+
        except Exception:
            pass


def main() -> int:
    _safe_console()
    log_path = os.path.join(DATA_DIR, "last_run_log.txt")
    sys.stdout = _Tee(sys.stdout, log_path)
    sys.stderr = sys.stdout        # 경고도 같은 자리에 모아 보이게

    ap = argparse.ArgumentParser(description="주식 테마 분류 및 테마 뉴스 수집기")
    ap.add_argument("--source", choices=["sample", "rss", "file"], default="sample")
    ap.add_argument("--input", help="--source file 일 때 읽을 텍스트 파일 경로")
    ap.add_argument("--hours", type=int, default=24, help="수집 시간 범위 (기본 24시간)")
    ap.add_argument("--market", choices=["ALL", "KR", "US"], default="ALL")
    ap.add_argument("--min-score", type=float, default=3.0, help="테마 분류 최소 점수")
    ap.add_argument("--xlsx", action="store_true", help="엑셀 리포트 생성")
    ap.add_argument("--no-inject", action="store_true", help="dashboard.html 갱신 생략")
    ap.add_argument("--notify", action="store_true",
                    help="속보를 텔레그램/슬랙/윈도우 알림으로 발송 (notify.json 설정 필요)")
    ap.add_argument("--search-limit", type=int, default=24, metavar="N",
                    help="구글뉴스 검색 요청 상한 (기본 24, 0이면 매체 RSS만 사용)")
    args = ap.parse_args()

    with open(os.path.join(BASE_DIR, "themes.json"), encoding="utf-8") as f:
        themes = json.load(f)["themes"]
    if args.market != "ALL":
        themes = [t for t in themes if t["market"] == args.market]

    print(f"■ 테마 {len(themes)}개 로드")
    print(f"■ 뉴스 수집 ({args.source}, 최근 {args.hours}시간)...")
    if args.source == "sample":
        raw = generate_sample(themes, args.hours)
    elif args.source == "file":
        if not args.input:
            ap.error("--source file 을 쓰려면 --input 경로가 필요합니다")
        raw = load_file(args.input)
    else:
        raw = collect_rss(themes, args.hours, args.search_limit)
    print(f"  → {len(raw)}건 수집")

    # 한 건도 못 받았는데 그대로 저장하면, 멀쩡히 보고 있던 화면이 빈 화면으로 덮인다.
    # 수집 실패는 '뉴스가 없다'가 아니라 '못 가져왔다'이므로 기존 데이터를 지키고 나간다.
    if not raw and args.source != "sample":
        print("\n  [!] 수집 결과가 0건이라 기존 데이터를 그대로 둡니다.")
        print("      화면의 뉴스가 안 바뀌는 건 이 때문입니다. 위의 실패 메시지를 확인해 주세요.")
        print("      회사망이라면: python collector.py --source rss --search-limit 0 로 매체 RSS만 시도해 보세요.")
        return 2

    classified = [classify(n, themes, args.min_score) for n in raw]
    classified.sort(key=lambda n: (-n["urgency"], n["published"]), reverse=False)
    classified.sort(key=lambda n: n["published"], reverse=True)
    matched = sum(1 for n in classified if n["themes"])
    print(f"  → 테마 분류 {matched}건 / 미분류 {len(classified) - matched}건")

    stats = build_theme_stats(classified, themes)
    payload = {
        "generated_at": datetime.now(KST).isoformat(),
        "window_hours": args.hours,
        "source": args.source,
        "themes": stats,
        "news": classified,
        "summary": {
            "total_news": len(classified),
            "classified": matched,
            "breaking": sum(1 for n in classified if n["urgency"] >= 3),
            "hot_theme": stats[0]["name"] if stats else None,
        },
    }

    print(f"■ 저장: {write_json(payload)}")
    if args.xlsx:
        p = write_xlsx(payload)
        if p:
            print(f"■ 엑셀: {p}")
    if args.notify:
        # 알림 실패가 수집 결과까지 날리지 않도록 여기서 막는다
        try:
            import notify
            sent = notify.run(payload, since_hours=args.hours)
            print(f"■ 속보 알림: {sent}건 발송")
        except Exception as exc:
            print(f"  [warn] 알림 발송 실패: {exc}", file=sys.stderr)

    if not args.no_inject:
        p = inject_into_dashboard(payload)
        if p:
            print(f"■ 대시보드 갱신: {p}")

    print("\n[상위 테마]")
    for t in stats[:8]:
        bar = "█" * int(min(t["heat"], 30) / 2)
        print(f"  {t['market']}  {t['name'][:26]:<28} {bar} {t['heat']:>5}  (뉴스 {t['news_count']}, 속보 {t['breaking_count']})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
